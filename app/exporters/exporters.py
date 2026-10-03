from __future__ import annotations
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import html
import json
from app.core.models import TranscriptResult, Segment


def _time_srt(seconds: float) -> str:
    ms = max(0, int(round(seconds * 1000)))
    h, rem = divmod(ms, 3600000)
    m, rem = divmod(rem, 60000)
    s, ms = divmod(rem, 1000)
    return f'{h:02}:{m:02}:{s:02},{ms:03}'


def _time_vtt(seconds: float) -> str:
    return _time_srt(seconds).replace(',', '.')


def _time_label(seconds: float) -> str:
    s = max(0, int(seconds))
    h, rem = divmod(s, 3600)
    m, sec = divmod(rem, 60)
    return f'{h:02d}:{m:02d}:{sec:02d}' if h else f'{m:02d}:{sec:02d}'


def _line(seg: Segment) -> str:
    who = f'{seg.speaker}：' if seg.speaker else ''
    return f'{who}{seg.text}'


def _write_vtt(result: TranscriptResult, path: Path, note: str | None = None) -> None:
    blocks = ['WEBVTT\n']
    if note:
        blocks.append(f'NOTE {note}\n')
    for s in result.segments:
        end = s.end if s.end > s.start else s.start + 2.0
        blocks.append(f'{_time_vtt(s.start)} --> {_time_vtt(end)}\n{_line(s)}\n')
    path.write_text('\n'.join(blocks), encoding='utf-8-sig')


def _review_html(result: TranscriptResult, title: str, audio_path: str | None, note: str | None) -> str:
    '''Build the offline KTV review player with importable text/subtitle tracks.'''
    audio_uri = ''
    if audio_path:
        try:
            audio_uri = Path(audio_path).resolve().as_uri()
        except Exception:
            audio_uri = ''

    segments = []
    for i, seg in enumerate(result.segments):
        end = seg.end if seg.end > seg.start else seg.start + 2.0
        segments.append({
            'i': i,
            'start': float(seg.start),
            'end': float(end),
            'text': seg.text,
            'speaker': seg.speaker or '',
            'label': _time_label(seg.start),
        })

    payload = json.dumps(segments, ensure_ascii=False).replace('</', '<\\/')
    ai_text = json.dumps(result.text or '', ensure_ascii=False).replace('</', '<\\/')
    note_html = f'<div class="note">{html.escape(note)}</div>' if note else ''
    safe_title = html.escape(title)
    safe_audio = html.escape(audio_uri, quote=True)

    page = r'''<!doctype html>
<html lang="zh-Hant-TW">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>__TITLE__｜錄音核對播放器</title>
<style>
:root{color-scheme:dark}
*{box-sizing:border-box}
body{font-family:"Microsoft JhengHei","Noto Sans TC",sans-serif;margin:0;background:#0b1220;color:#e5e7eb}
header{position:sticky;top:0;background:rgba(11,18,32,.97);padding:14px 20px;border-bottom:1px solid #334155;z-index:10;backdrop-filter:blur(8px)}
h1{font-size:20px;margin:0 0 10px}.tools{display:flex;gap:8px;flex-wrap:wrap;align-items:center}.row{display:flex;gap:8px;flex-wrap:wrap;align-items:center;margin-top:8px}
audio{width:min(880px,100%);height:40px}input[type=file]{max-width:330px}
select,input[type=search],textarea{padding:8px 10px;border-radius:7px;border:1px solid #475569;background:#111827;color:#fff}
select{min-width:210px}#search{width:min(360px,80vw)}
button{padding:7px 10px;border-radius:7px;border:1px solid #475569;background:#1e293b;color:#fff;cursor:pointer}button:hover{background:#334155}
main{max-width:1180px;margin:auto;padding:16px 20px 80px}.note{background:#49380b;border:1px solid #8a6d1d;padding:10px;border-radius:8px;margin:10px 0}.small{font-size:12px;color:#94a3b8}
.panel{background:#0f172a;border:1px solid #334155;border-radius:10px;padding:12px;margin:12px 0}.panel-title{font-weight:700;margin-bottom:8px;color:#bae6fd}
#fulltext{max-height:36vh;overflow:auto;line-height:2.15;font-size:1.08rem;padding:12px;border-radius:8px;background:#0a1324;border:1px solid #25354d;scroll-behavior:smooth}
.phrase{display:inline;padding:3px 5px;margin:1px 1px;border-radius:5px;cursor:pointer;transition:background .12s,color .12s}.phrase:hover{background:#1e293b}.phrase.active{color:#fff;background:linear-gradient(90deg,#b45309 var(--prog,0%),#164e63 var(--prog,0%));box-shadow:0 0 0 1px #38bdf8 inset}.phrase.speaker:before{content:attr(data-speaker) '：';color:#fcd34d;font-weight:700}
#now{margin-top:8px;color:#7dd3fc;font-variant-numeric:tabular-nums}.badge{display:inline-block;padding:2px 7px;border-radius:999px;font-size:12px;margin-left:6px}.exact{background:#064e3b;color:#a7f3d0}.estimate{background:#713f12;color:#fde68a}
.seg{display:grid;grid-template-columns:92px 1fr;gap:12px;padding:10px 8px;border-bottom:1px solid #263449;cursor:pointer;border-radius:6px}.seg:hover{background:#172033}.seg.active{background:#123b5d;outline:1px solid #38bdf8}.time{color:#7dd3fc;font-variant-numeric:tabular-nums}.speakerTag{font-weight:700;color:#fcd34d;margin-right:6px}.text{line-height:1.65}.alt{margin-top:5px;padding-top:5px;border-top:1px dashed #3b4b63;color:#cbd5e1}.alt:before{content:'對照：';color:#a5b4fc;font-weight:700}
#pastePanel{display:none;margin-top:8px}#pasteText{width:100%;min-height:120px;resize:vertical}.status{color:#a5f3fc;margin-top:7px}.warn{color:#fde68a}.ok{color:#86efac}
@media(max-width:700px){header{position:static}.seg{grid-template-columns:72px 1fr}#fulltext{max-height:44vh}}
</style>
</head>
<body>
<header>
<h1>__TITLE__｜錄音與逐字稿核對｜KTV 同步</h1>
<div class="tools">
<audio id="audio" controls preload="metadata" src="__AUDIO__"></audio>
<label class="small">音檔無法載入時：<input id="pickAudio" type="file" accept="audio/*,video/*"></label>
</div>
<div class="row">
<label>全文文字軌：<select id="track"></select></label>
<input id="search" type="search" placeholder="搜尋目前文字軌…">
<button id="follow">自動跟隨：開</button>
<button id="dual">下方雙稿：開</button>
</div>
<div class="row">
<label class="small">加入其他翻譯／逐字稿：<input id="pickText" type="file" accept=".txt,.srt,.vtt,.json,.md,.markdown,text/plain,text/markdown,text/vtt,application/json"></label>
<button id="pasteBtn">貼上文字</button>
<button id="downloadVtt">匯出目前文字軌 VTT</button>
<span id="trackStatus" class="status"></span>
</div>
<div id="pastePanel"><textarea id="pasteText" placeholder="可直接從 Word、記事本或其他翻譯結果貼上文字。沒有時間碼時會依原始時間軸順序估算對齊。"></textarea><div class="row"><button id="applyPaste">套用貼上文字</button><button id="cancelPaste">取消</button></div></div>
<div class="small">SRT/VTT 會保留原本時間碼；TXT/Markdown／貼上文字會依原始逐字稿段落順序自動估算時間。點上方全文或下方段落都可跳到錄音位置。</div>
</header>
<main>
__NOTE__
<section class="panel"><div class="panel-title">完整全文 KTV 同步 <span id="quality"></span></div><div id="fulltext"></div><div id="now">待播放</div></section>
<section class="panel"><div class="panel-title">時間軸核對</div><div id="list"></div></section>
</main>
<script>
const baseSegs=__PAYLOAD__;
const aiText=__AI_TEXT__;
const audio=document.getElementById('audio'), list=document.getElementById('list'), full=document.getElementById('fulltext');
const trackSel=document.getElementById('track'), statusEl=document.getElementById('trackStatus'), qualityEl=document.getElementById('quality'), nowEl=document.getElementById('now');
let follow=true, dual=true, active=-1, currentKey='base', importedSeq=0;
const tracks={base:{name:'原始時間軸逐字稿',segs:baseSegs,exact:true,source:'內建'}};

function esc(s){return String(s??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));}
function parseTs(v){v=String(v).trim().replace(',','.');const p=v.split(':').map(Number);if(p.length===3)return p[0]*3600+p[1]*60+p[2];if(p.length===2)return p[0]*60+p[1];return Number(v)||0;}
function timeLabel(sec){sec=Math.max(0,Math.floor(Number(sec)||0));const h=Math.floor(sec/3600),m=Math.floor((sec%3600)/60),s=sec%60;return h?`${String(h).padStart(2,'0')}:${String(m).padStart(2,'0')}:${String(s).padStart(2,'0')}`:`${String(m).padStart(2,'0')}:${String(s).padStart(2,'0')}`;}
function splitPlainText(text){
  const lines=String(text||'').replace(/\r/g,'').split('\n').map(x=>x.trim()).filter(Boolean);
  if(lines.length>1)return lines;
  const one=lines[0]||'';
  const parts=one.match(/[^。！？!?；;]+[。！？!?；;]?/g)||[one];
  return parts.map(x=>x.trim()).filter(Boolean);
}
function alignPlainText(text){
  const parts=splitPlainText(text); if(!parts.length)return [];
  if(!baseSegs.length)return parts.map((t,i)=>({i,start:i*2,end:(i+1)*2,text:t,speaker:'',label:timeLabel(i*2)}));
  const n=baseSegs.length, m=parts.length, out=[];
  for(let i=0;i<m;i++){
    let a=Math.floor(i*n/m), b=Math.max(a,Math.ceil((i+1)*n/m)-1); a=Math.min(a,n-1); b=Math.min(b,n-1);
    out.push({i,start:Number(baseSegs[a].start)||0,end:Number(baseSegs[b].end)||Number(baseSegs[b].start)+2,text:parts[i],speaker:'',label:timeLabel(baseSegs[a].start)});
  }
  return out;
}
function parseTimedText(text){
  const src=String(text||'').replace(/\r/g,''); const lines=src.split('\n'); const out=[];
  for(let i=0;i<lines.length;i++){
    const m=lines[i].match(/((?:\d{1,2}:)?\d{2}:\d{2}[.,]\d{1,3})\s*-->\s*((?:\d{1,2}:)?\d{2}:\d{2}[.,]\d{1,3})/);
    if(!m)continue; const start=parseTs(m[1]), end=parseTs(m[2]); const body=[]; i++;
    while(i<lines.length && lines[i].trim() && !/-->/.test(lines[i])){body.push(lines[i].trim());i++;}
    let txt=body.join(' ').replace(/^\d+\s*$/,'').trim(); if(!txt)continue;
    let speaker=''; const sm=txt.match(/^([^：:]{1,30})[：:]\s*(.+)$/); if(sm){speaker=sm[1];txt=sm[2];}
    out.push({i:out.length,start,end:end>start?end:start+2,text:txt,speaker,label:timeLabel(start)});
  }
  return out;
}
function parseJsonTrack(text){
  const data=JSON.parse(text); const arr=Array.isArray(data)?data:(Array.isArray(data.segments)?data.segments:[]);
  return arr.map((s,i)=>({i,start:Number(s.start)||0,end:Number(s.end)>Number(s.start)?Number(s.end):(Number(s.start)||0)+2,text:String(s.text||''),speaker:String(s.speaker||''),label:timeLabel(s.start)})).filter(s=>s.text);
}
function addTrack(name,segs,exact,source){
  const key='import'+(++importedSeq); tracks[key]={name,segs,exact,source}; currentKey=key; refreshTrackSelect(); renderAll();
  statusEl.className='status '+(exact?'ok':'warn'); statusEl.textContent=exact?'已載入含時間碼文字軌。':'已載入文字；時間為依原始逐字稿順序估算。';
}
function refreshTrackSelect(){
  trackSel.innerHTML=''; Object.entries(tracks).forEach(([k,t])=>{const o=document.createElement('option');o.value=k;o.textContent=t.name;trackSel.appendChild(o);});trackSel.value=currentKey;
}
if(aiText.trim()){tracks.ai={name:'AI 整理全文（估算同步）',segs:alignPlainText(aiText),exact:false,source:'AI整理'};}
function currentTrack(){return tracks[currentKey]||tracks.base;}
function findIdx(segs,t){for(let i=0;i<segs.length;i++){if(t>=segs[i].start&&t<segs[i].end)return i;}return -1;}
function findNearestByTime(segs,t){let best=-1,dist=Infinity;segs.forEach((s,i)=>{const d=Math.abs((s.start+s.end)/2-t);if(d<dist){dist=d;best=i;}});return best;}
function renderFull(){
  const t=currentTrack(); full.innerHTML=''; qualityEl.innerHTML=t.exact?'<span class="badge exact">精準時間碼</span>':'<span class="badge estimate">估算對齊</span>';
  t.segs.forEach((s,idx)=>{const sp=document.createElement('span');sp.className='phrase'+(s.speaker?' speaker':'');sp.dataset.idx=idx;sp.dataset.speaker=s.speaker||'';sp.textContent=s.text+' ';sp.title=`${timeLabel(s.start)}–${timeLabel(s.end)}`;sp.onclick=()=>{audio.currentTime=s.start;audio.play().catch(()=>{});setActive(idx,true);};full.appendChild(sp);});
}
function altAt(baseSeg){
  const t=currentTrack(); if(currentKey==='base'||!dual)return null; const mid=(baseSeg.start+baseSeg.end)/2; let idx=findIdx(t.segs,mid); if(idx<0)idx=findNearestByTime(t.segs,mid); return idx>=0?t.segs[idx]:null;
}
function renderList(filter=''){
  const q=filter.trim().toLowerCase(); list.innerHTML=''; const source=baseSegs;
  source.forEach((s,idx)=>{const alt=altAt(s);const hay=(s.text+' '+s.speaker+' '+(alt?alt.text:'')).toLowerCase();if(q&&!hay.includes(q))return;
    const d=document.createElement('div');d.className='seg';d.dataset.idx=idx;
    d.innerHTML=`<div class="time">${esc(timeLabel(s.start))}</div><div class="text">${s.speaker?'<span class="speakerTag">'+esc(s.speaker)+'：</span>':''}${esc(s.text)}${alt?'<div class="alt">'+esc(alt.text)+'</div>':''}</div>`;
    d.onclick=()=>{audio.currentTime=s.start;audio.play().catch(()=>{});updateFromTime();};list.appendChild(d);
  });
}
function renderAll(){renderFull();renderList(document.getElementById('search').value||'');active=-1;updateFromTime();}
function setActive(idx,scroll=false){
  const t=currentTrack(); if(idx<0||idx>=t.segs.length)return; active=idx;
  document.querySelectorAll('.phrase.active').forEach(x=>{x.classList.remove('active');x.style.removeProperty('--prog');}); const ph=document.querySelector(`.phrase[data-idx="${idx}"]`);if(ph){ph.classList.add('active');if(scroll&&follow)ph.scrollIntoView({block:'center',behavior:'smooth'});}
  const s=t.segs[idx];nowEl.textContent=`${timeLabel(s.start)}–${timeLabel(s.end)}｜${s.speaker?s.speaker+'｜':''}${s.text}`;
  document.querySelectorAll('.seg.active').forEach(x=>x.classList.remove('active')); const bi=findIdx(baseSegs,audio.currentTime);const el=document.querySelector(`.seg[data-idx="${bi>=0?bi:0}"]`);if(el){el.classList.add('active');}
}
function updateFromTime(){
  const t=currentTrack(), now=audio.currentTime||0; let idx=findIdx(t.segs,now);if(idx<0)return;setActive(idx,true);const s=t.segs[idx];const dur=Math.max(.05,s.end-s.start);const p=Math.max(0,Math.min(100,(now-s.start)/dur*100));const ph=document.querySelector(`.phrase[data-idx="${idx}"]`);if(ph)ph.style.setProperty('--prog',p+'%');
}
function toVtt(segs){let out='WEBVTT\n\n';segs.forEach(s=>{const fmt=x=>{const ms=Math.max(0,Math.round(x*1000)),h=Math.floor(ms/3600000),m=Math.floor(ms%3600000/60000),sec=Math.floor(ms%60000/1000),r=ms%1000;return `${String(h).padStart(2,'0')}:${String(m).padStart(2,'0')}:${String(sec).padStart(2,'0')}.${String(r).padStart(3,'0')}`;};out+=`${fmt(s.start)} --> ${fmt(s.end)}\n${s.speaker?s.speaker+'：':''}${s.text}\n\n`;});return out;}
trackSel.onchange=()=>{currentKey=trackSel.value;renderAll();};audio.addEventListener('timeupdate',updateFromTime);audio.addEventListener('seeked',updateFromTime);
document.getElementById('search').addEventListener('input',e=>renderList(e.target.value));
document.getElementById('pickAudio').addEventListener('change',e=>{const f=e.target.files[0];if(f){audio.src=URL.createObjectURL(f);audio.load();}});
document.getElementById('follow').onclick=e=>{follow=!follow;e.target.textContent='自動跟隨：'+(follow?'開':'關');};
document.getElementById('dual').onclick=e=>{dual=!dual;e.target.textContent='下方雙稿：'+(dual?'開':'關');renderList(document.getElementById('search').value||'');};
document.getElementById('pasteBtn').onclick=()=>{document.getElementById('pastePanel').style.display='block';document.getElementById('pasteText').focus();};
document.getElementById('cancelPaste').onclick=()=>{document.getElementById('pastePanel').style.display='none';};
document.getElementById('applyPaste').onclick=()=>{const x=document.getElementById('pasteText').value.trim();if(!x)return;addTrack('貼上文字',alignPlainText(x),false,'貼上');document.getElementById('pastePanel').style.display='none';};
document.getElementById('pickText').addEventListener('change',async e=>{const f=e.target.files[0];if(!f)return;try{const text=await f.text();let segs=[],exact=false;const n=f.name.toLowerCase();if(n.endsWith('.srt')||n.endsWith('.vtt')){segs=parseTimedText(text);exact=segs.length>0;}else if(n.endsWith('.json')){segs=parseJsonTrack(text);exact=segs.length>0;}else{segs=alignPlainText(text);}if(!segs.length)throw new Error('找不到可用文字內容');addTrack(f.name,segs,exact,f.name);}catch(err){statusEl.className='status warn';statusEl.textContent='文字檔載入失敗：'+err.message;}});
document.getElementById('downloadVtt').onclick=()=>{const t=currentTrack();if(!t.segs.length)return;const blob=new Blob([toVtt(t.segs)],{type:'text/vtt;charset=utf-8'});const a=document.createElement('a');a.href=URL.createObjectURL(blob);a.download=(t.name||'transcript').replace(/[\\/:*?"<>|]/g,'_')+'.vtt';a.click();setTimeout(()=>URL.revokeObjectURL(a.href),1000);};
refreshTrackSelect();renderAll();
</script>
</body></html>'''
    return (page
            .replace('__TITLE__', safe_title)
            .replace('__AUDIO__', safe_audio)
            .replace('__NOTE__', note_html)
            .replace('__PAYLOAD__', payload)
            .replace('__AI_TEXT__', ai_text))

def export_all(result: TranscriptResult, base_path: str, formats: list[str], note: str | None = None,
               source_audio: str | None = None) -> list[str]:
    """Export transcript formats plus an optional clickable audio review player."""
    base = Path(base_path)
    base.parent.mkdir(parents=True, exist_ok=True)
    out = []
    formats = list(dict.fromkeys(formats or []))
    if 'txt' in formats:
        p = base.with_suffix('.txt')
        body = result.text or '\n'.join(_line(x) for x in result.segments)
        if note:
            body = note + '\n\n' + body
        p.write_text(body, encoding='utf-8-sig')
        out.append(str(p))
    if 'md' in formats:
        p = base.with_suffix('.md')
        lines = [f'# {base.stem}', '', f'- 辨識引擎：{result.engine}']
        if result.language:
            lines.append(f'- 語言：{result.language}')
        if source_audio:
            lines.append(f'- 來源音檔：{Path(source_audio).name}')
        if note:
            lines += ['', f'> {note}']
        lines += ['', '## 時間軸逐字稿', '']
        if result.segments:
            for s in result.segments:
                end = s.end if s.end > s.start else s.start + 2.0
                speaker = f'{s.speaker}：' if s.speaker else ''
                lines.append(f'- **[{_time_label(s.start)}–{_time_label(end)}]** {speaker}{s.text}')
        elif result.text:
            lines.append(result.text)
        if result.segments and result.text:
            timeline_text = '\n'.join(_line(x) for x in result.segments).strip()
            cleaned = result.text.strip()
            if cleaned and cleaned != timeline_text:
                lines += ['', '## AI 整理後全文（參考）', '', cleaned]
        p.write_text('\n'.join(lines) + '\n', encoding='utf-8-sig')
        out.append(str(p))
    if 'srt' in formats and result.segments:
        p = base.with_suffix('.srt')
        blocks = []
        for i, s in enumerate(result.segments, 1):
            end = s.end if s.end > s.start else s.start + 2.0
            blocks.append(f'{i}\n{_time_srt(s.start)} --> {_time_srt(end)}\n{_line(s)}')
        p.write_text('\n\n'.join(blocks), encoding='utf-8-sig')
        out.append(str(p))
    if 'vtt' in formats and result.segments:
        p = base.with_suffix('.vtt')
        _write_vtt(result, p, note)
        out.append(str(p))
    if 'docx' in formats:
        p = base.with_suffix('.docx')
        doc = Document()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)
        doc.add_heading(base.stem, level=1)
        doc.add_paragraph(f'辨識引擎：{result.engine}')
        if result.language:
            doc.add_paragraph(f'語言：{result.language}')
        if source_audio:
            doc.add_paragraph(f'來源音檔：{Path(source_audio).name}')
        if note:
            doc.add_paragraph(note)

        # v3.7: Word 的主體永遠先放「時間軸逐字稿」。
        # 先前混合模式把 Gemini 整理後全文放在最前面，長錄音看起來
        # 會像一大坨文字，使用者很容易以為時間資訊遺失。
        if result.segments:
            doc.add_heading('時間軸逐字稿（主要核對版）', level=2)
            doc.add_paragraph('每列皆保留錄音開始／結束時間；可搭配同名「_錄音核對.html」或「_字幕.vtt」核對原音。')
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Table Grid'
            table.autofit = False
            hdr = table.rows[0].cells
            hdr[0].text = '時間'
            hdr[1].text = '講者'
            hdr[2].text = '逐字內容'
            # 讓時間／講者欄保持精簡，把大部分寬度留給逐字內容。
            widths = (Inches(1.45), Inches(0.70), Inches(8.10))
            for cell, width in zip(hdr, widths):
                cell.width = width
            # 多頁表格重複顯示標題列，方便長錄音核對。
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            tbl_header = OxmlElement('w:tblHeader')
            tbl_header.set(qn('w:val'), 'true')
            tr_pr.append(tbl_header)
            for s in result.segments:
                end = s.end if s.end > s.start else s.start + 2.0
                row = table.add_row().cells
                for cell, width in zip(row, widths):
                    cell.width = width
                row[0].text = f'{_time_label(s.start)}–{_time_label(end)}'
                row[1].text = s.speaker or ''
                row[2].text = s.text
                for cell in (row[0], row[1]):
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(9)
                for para in row[2].paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(11)

            # 混合模式／Gemini 整理文字是附錄，不取代時間軸。
            if result.text:
                timeline_text = '\n'.join(_line(x) for x in result.segments).strip()
                cleaned = result.text.strip()
                if cleaned and cleaned != timeline_text:
                    doc.add_page_break()
                    doc.add_heading('AI 整理後全文（參考附錄）', level=2)
                    doc.add_paragraph('此區為閱讀版文字；錄音時間核對請以前面的「時間軸逐字稿」為準。')
                    doc.add_paragraph(cleaned)
        else:
            doc.add_heading('逐字稿', level=2)
            doc.add_paragraph('本次結果未包含可用的時間軸資訊。')
            if result.text:
                doc.add_paragraph(result.text)
        doc.save(p)
        out.append(str(p))
    if 'review' in formats and result.segments:
        # Review mode always exports a WebVTT subtitle alongside the HTML player.
        vtt = base.with_name(base.stem + '_字幕').with_suffix('.vtt')
        _write_vtt(result, vtt, note)
        if str(vtt) not in out:
            out.append(str(vtt))
        page = base.with_name(base.stem + '_錄音核對').with_suffix('.html')
        page.write_text(_review_html(result, base.stem, source_audio, note), encoding='utf-8')
        out.append(str(page))
    return out
