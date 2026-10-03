from pathlib import Path
from docx import Document
from app.core.models import TranscriptResult, Segment
from app.exporters.exporters import export_all


def test_word_timeline_is_primary_and_cleaned_text_is_appendix(tmp_path: Path):
    r = TranscriptResult(
        text='這是 Gemini 整理後的一大段閱讀版全文。',
        segments=[
            Segment(0.0, 3.2, '大家早安。', '講者1'),
            Segment(3.2, 7.8, '今天先確認進度。', '講者2'),
        ],
        language='zh',
        engine='混合模式',
    )
    out = export_all(r, str(tmp_path/'meeting_逐字稿'), ['docx'], source_audio='meeting.m4a')
    docx = Path(out[0])
    doc = Document(docx)
    paras = [p.text for p in doc.paragraphs]
    assert '時間軸逐字稿（主要核對版）' in paras
    assert 'AI 整理後全文（參考附錄）' in paras
    assert paras.index('時間軸逐字稿（主要核對版）') < paras.index('AI 整理後全文（參考附錄）')
    assert len(doc.tables) == 1
    rows = doc.tables[0].rows
    assert [c.text for c in rows[0].cells] == ['時間', '講者', '逐字內容']
    assert rows[1].cells[0].text == '00:00–00:03'
    assert rows[1].cells[1].text == '講者1'
    assert rows[1].cells[2].text == '大家早安。'
    assert rows[2].cells[0].text == '00:03–00:07'
